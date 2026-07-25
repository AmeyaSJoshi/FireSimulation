#!/usr/bin/env python3
"""Build the compact browser/server parameter table from the published FCCS workbook."""

from __future__ import annotations

import json
import sys
from pathlib import Path

from openpyxl import load_workbook


INPUT = Path(sys.argv[1] if len(sys.argv) > 1 else "/private/tmp/Global_fuelbeds_parameters_v1.2.xlsx")
OUTPUT = Path(sys.argv[2] if len(sys.argv) > 2 else "public/global-fuelbed-parameters.json")

FIELD_MAP = {
    "fuelbed": "FUELBED",
    "joinValue": "JOIN_VALUE",
    "biome": "Biome",
    "landCover": "LandCover",
    "treeCoverPercent": "Tree Cover (%)",
    "treeOverstoryCoverPercent": "TO_Cover (%)",
    "treeMidstoryCoverPercent": "TM_Cover (%)",
    "treeOverstoryHeightMeters": "TO_Height (m)",
    "treeOverstoryLiveCrownBaseMeters": "TO_HLC (m)",
    "treeMidstoryHeightMeters": "TM_Height (m)",
    "treeMidstoryLiveCrownBaseMeters": "TM_HLC (m)",
    "treeLadderFuelPresent": "T_Ladder",
    "treeLadderFuelType": "T_LadderType",
    "shrubCoverPercent": "Shrub Cover (%)",
    "grassCoverPercent": "Grass Cover (%)",
    "grassLivePercent": "G_live (%)",
    "shrubLivePercent": "S_Live (%)",
    "woodyCoverPercent": "Woody Cover (%)",
    "grassHeightMeters": "G_height (m)",
    "shrubHeightMeters": "S_Height (m)",
    "woodyDepthCm": "W_depth (cm)",
    "duffDepthInches": "DU_Depth (in)",
    "grassLoadMgPerHa": "G_Load (Mg/ha)",
    "dead1hLoadMgPerHa": "W_1hLoad (Mg/ha)",
    "dead10hLoadMgPerHa": "W_10h Load (Mg/ha)",
    "dead100hLoadMgPerHa": "W_100h Load (Mg/ha)",
    "dead1000hLoadMgPerHa": "W_1000h Load (Mg/ha)",
    "litterCoverPercent": "Litter_Cover (%)",
    "litterDepthCm": "L_depth (cm)",
    "litterArrangement": "Litter_Arr",
    "upperDuffCoverPercent": "DU_cover (%)",
    "lowerDuffCoverPercent": "DL_cover (%)",
}


def json_value(value):
    if value is None:
        return None
    if isinstance(value, (int, float, str, bool)):
        return value
    return str(value)


workbook = load_workbook(INPUT, read_only=True, data_only=True)
sheet = workbook["Fuelbeds_metric"]
rows = sheet.iter_rows(values_only=True)
headers = next(rows)
index = {header: position for position, header in enumerate(headers) if header}

compact_rows = []
for source_row in rows:
    fuelbed = source_row[index["FUELBED"]]
    join_value = source_row[index["JOIN_VALUE"]]
    if fuelbed is None or join_value is None:
        continue
    record = {}
    for output_key, source_key in FIELD_MAP.items():
        value = source_row[index[source_key]]
        if isinstance(value, float) and value.is_integer():
            value = int(value)
        record[output_key] = json_value(value)
    compact_rows.append(record)

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
OUTPUT.write_text(json.dumps({
    "source": {
        "title": "Pettinari and Chuvieco Global Fuelbed Dataset",
        "parameterVersion": "v1.2",
        "doi": "10.1594/PANGAEA.849808",
        "license": "CC-BY-NC-SA-3.0"
    },
    "rows": compact_rows
}, separators=(",", ":")) + "\n", encoding="utf-8")
print(f"wrote {len(compact_rows)} rows to {OUTPUT}")
